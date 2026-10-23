"""Desktop prices must be compared only with stated matching components."""

from openpyxl import load_workbook

import pc_prices
import profiles
import telegram_notify


def pc(item_id: int, price: int, *, title: str, description: str = "") -> dict:
    return {
        "id": item_id, "type": "item", "title": title,
        "description": description,
        "urlPath": f"/kazan/nastolnye_kompyutery/pc_{item_id}",
        "priceDetailed": {"value": price},
        "closedItemsText": "3 завершённых объявления",
        "contacts": {"phone": True},
    }


def test_pc_group_requires_cpu_and_ram_and_separates_graphics():
    first = pc(1, 7000, title="ПК i5-4460, 8 ГБ ОЗУ, GTX 750 Ti")
    second = pc(2, 7500, title="Компьютер", description="Intel Core i5 4570, RAM 8 GB, GeForce GTX 750 Ti")

    assert pc_prices.group_key(first) == pc_prices.group_key(second)
    assert pc_prices.group_key(pc(3, 5000, title="Системный блок i5 8gb")) is None
    assert pc_prices.group_key(pc(4, 5000, title="Ноутбук i5-4460, 8 ГБ ОЗУ, GTX 750 Ti")) is None
    assert pc_prices.group_key(pc(5, 5000, title="ПК i5-750, 8 ГБ ОЗУ, GTX 750 Ti")) is None
    assert pc_prices.group_key(pc(6, 5000, title="ПК i3 2100, 8gb, 500gb, GTX 650")) is not None
    assert pc_prices.group_key(pc(7, 7000, title="ПК Ryzen 5 5600G, 16 ГБ ОЗУ", description="Встроенная графика")) is not None
    unknown = pc_prices.group_key(pc(8, 7000, title="ПК i5-4460, 8 ГБ ОЗУ"))
    assert "видеокарта не указана" in unknown
    assert unknown != pc_prices.group_key(first)


def test_pc_reference_separates_graphics_bundle_and_incomplete(tmp_path):
    base = {
        str(i): pc(i, price, title="ПК i5-4460, 8 ГБ ОЗУ, GTX 750 Ti", description="Рабочий")
        for i, price in enumerate((7000, 7500, 8000), 1)
    }
    base["4"] = pc(4, 2000, title="ПК i5-4460, 8 ГБ ОЗУ, GTX 750 Ti", description="На запчасти")
    base["5"] = pc(5, 9000, title="ПК i5-4460, 8 ГБ ОЗУ, GTX 750 Ti с монитором")
    base["6"] = pc(6, 4000, title="ПК i5-4460, 8 ГБ ОЗУ, GT 730")
    profile = profiles.PCPricesProfile(base)

    cheap = profile.alert(pc(7, 6000, title="ПК i5-4570, 8 ГБ ОЗУ, GTX 750 Ti"))
    assert cheap is not None
    assert cheap["market"] == "7 500 ₽"
    assert cheap["vs_market"] == "-20%"
    assert "🖥️" in telegram_notify.message_for("pc-prices", "pcs", cheap)
    assert "на 20% дешевле медианы" in telegram_notify.message_for("pc-prices", "pcs", cheap)
    assert profile.alert(pc(8, 6000, title="ПК i5-4570, 8 ГБ ОЗУ, GT 730"))["vs_market"] == ""
    assert profile.alert(pc(9, 6000, title="Системный блок"))["price_assessment"] == "нет достаточной выборки"

    target = tmp_path / "pcs.xlsx"
    counts = profile.write_table(target, {})
    book = load_workbook(target)
    assert counts["групп с рыночной ценой"] == 1
    assert book["Рынок"].max_row == 2
    assert any(row[3] == 7500 for row in book["Объявления"].iter_rows(min_row=2, values_only=True))


def test_pc_without_stated_gpu_gets_cautious_estimate():
    base = {str(i): pc(i, price, title="ПК i3-2100, 8 ГБ ОЗУ")
            for i, price in enumerate((4000, 4500, 5000), 1)}
    profile = profiles.PCPricesProfile(base)

    alert = profile.alert(pc(10, 3600, title="Системный блок i3-2120, 8 ГБ ОЗУ"))
    assert alert is not None
    assert alert["market"] == "4 500 ₽"
    assert alert["price_assessment"].startswith("ориентировочно")
    assert any("видеокарта не указана" in note for note in alert["notes"])


def test_incomplete_and_broken_pcs_do_not_alert():
    profile = profiles.PCPricesProfile({})
    for description in ("Без процессора", "На запчасти", "Не включается", "Без накопителя"):
        assert profile.alert(pc(1, 2000, title="Системный блок", description=description)) is None
    assert profile.alert(pc(2, 7000, title="Системный блок", description="Полностью рабочий")) is not None
